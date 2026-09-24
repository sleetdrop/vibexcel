"""Read-only independent audit of saved Sudoku candidate displays.

Run with the bundled Python interpreter after Excel has saved the test copy.
The chain may remove candidates by sound deductions, so its set must be a
subset of the basic row/column/box set; equality is not expected.
"""

import argparse
import sys
import warnings

import openpyxl


PAGES = (
    ("01 Easy", "_Engine"),
    ("02 Medium", "_Engine Medium"),
    ("03 Hard", "_Engine Hard"),
    ("04 Expert", "_Engine Expert"),
    ("05 Master", "_Engine Master"),
)
CANDIDATE_OFFSETS = ((-1, -1), (-1, 0), (-1, 1), (0, -1), (0, 1), (1, -1), (1, 0), (1, 1))


def digits(value):
    return {int(char) for char in str(value or "") if char in "123456789"}


def basic_candidates(board, row, col):
    used = set(board[row])
    used.update(board[r][col] for r in range(9))
    used.update(board[r][c] for r in range(row // 3 * 3, row // 3 * 3 + 3)
                for c in range(col // 3 * 3, col // 3 * 3 + 3))
    return set(range(1, 10)) - used


def audit(path):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Conditional Formatting extension")
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    failures = []
    summary = []
    for page_name, engine_name in PAGES:
        page, engine = workbook[page_name], workbook[engine_name]
        board = [[engine.cell(3 + r, 2 + c).value or 0 for c in range(9)] for r in range(9)]
        visible_count = refined_count = 0
        for row in range(9):
            for col in range(9):
                center_row, center_col = 5 + 3 * row, 3 + 4 * col
                center = page.cell(center_row, center_col).value or 0
                live = board[row][col]
                label = f"{page_name} R{row + 1}C{col + 1}"
                if center != live:
                    failures.append(f"{label}: center {center!r} differs from engine {live!r}")
                refined = digits(engine.cell(15 + row, 2 + col).value)
                chain = digits(engine.cell(102 + row, 12 + col).value)
                shown = set()
                for dr, dc in CANDIDATE_OFFSETS:
                    shown.update(digits(page.cell(center_row + dr, center_col + dc).value))
                if refined != chain:
                    failures.append(f"{label}: projection {sorted(refined)} differs from chain {sorted(chain)}")
                if live:
                    if refined or shown:
                        failures.append(f"{label}: filled square shows candidates {sorted(refined)} / {sorted(shown)}")
                    continue
                legal = basic_candidates(board, row, col)
                if not refined <= legal:
                    failures.append(f"{label}: illegal refined digits {sorted(refined - legal)}")
                if shown != refined:
                    failures.append(f"{label}: shown {sorted(shown)} differs from engine {sorted(refined)}")
                refined_count += len(refined)
                visible_count += len(shown)
        summary.append((page_name, page["AM5"].value, page["AU8"].value, refined_count, visible_count))
    workbook.close()
    for name, status, conflicts, refined_count, visible_count in summary:
        print(f"{name}: status={status!r} conflicts={conflicts} chain_digits={refined_count} shown_digits={visible_count}")
    if failures:
        for failure in failures:
            print(f"FAIL {failure}", file=sys.stderr)
    else:
        print("PASS: all 405 squares match the saved engine projections and basic legality constraints")
    return not failures


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", help="Excel-saved .xlsx file to inspect without modifying it")
    args = parser.parse_args()
    sys.exit(0 if audit(args.workbook) else 1)
